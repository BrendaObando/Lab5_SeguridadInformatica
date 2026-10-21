from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt

root = Path(__file__).resolve().parents[1]
source = Path(r'C:\Users\Bren\Documents\Laboratorio5_SeguridadInformatica.docx')
destination = root / 'Laboratorio5_SeguridadInformatica_Evidencias.docx'
evidence = root / 'evidencias'

document = Document(source)

def heading(text, size):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    return paragraph

document.add_page_break()
heading('Evidencias — Laboratorio 5', 16)
document.add_paragraph('Proyecto: UNA Chat. Se implementó la validación segura de URLs y la prevención de inyección de scripts. Las pruebas se ejecutaron localmente el 16 de octubre de 2026.')

def add_result(title, image_name, result_file, explanation):
    heading(title, 13)
    document.add_paragraph(explanation)
    image = evidence / image_name
    if image.exists():
        document.add_picture(str(image), width=Inches(6.3))
    output = (evidence / result_file).read_text(encoding='utf-8')
    paragraph = document.add_paragraph()
    run = paragraph.add_run(output)
    run.font.name = 'Consolas'
    run.font.size = Pt(8)

add_result('Evidencia 1 — Inyección de script bloqueada', 'evidencia-1-inyeccion-bloqueada.jpg', 'resultado-unitarias.txt', 'El payload <script>alert(...)</script> se recibe como texto; no se crea una etiqueta script ni se abre un diálogo. La prueba unitaria correspondiente figura en el resultado adjunto.')
add_result('Evidencia 2 — Pruebas unitarias', 'funcionalidad-url-segura.jpg', 'resultado-unitarias.txt', 'Mocha valida teléfono, imágenes, videos de YouTube, clasificación de URLs y la protección contra scripts.')
add_result('Evidencia 3 — Pruebas de integración', 'funcionalidad-url-segura.jpg', 'resultado-integracion.txt', 'Jest y socket.io-client prueban el evento real del backend y la librería una-lib para imagen, video y script.')
add_result('Evidencia 4 — Pruebas end-to-end', 'funcionalidad-url-segura.jpg', 'resultado-e2e.txt', 'Playwright recorre el navegador real con Chrome: muestra imagen, reproductor de video y bloquea el alert de script.')
heading('Cambios implementados', 13)
document.add_paragraph('• Se sustituyó la inserción de HTML de usuario por nodos DOM y textContent.\n• Las imágenes se limitan a URLs HTTP(S) con extensiones permitidas.\n• Los videos se limitan a enlaces válidos de YouTube y se convierten a una URL embed controlada.\n• Se añadieron suites Mocha, Jest/Socket.IO y Playwright.')
document.save(destination)
print(destination)
